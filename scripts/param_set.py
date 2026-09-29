# -*- coding: utf-8 -*-
"""wsp 參數寫入工具：依彙整主表的「所有參數設定值」（最佳參數）產生 .wsp 副本

產出：<原檔名>_BEST.wsp（param_set 寫入 Value 與 LocalCurrentValue2 後存成副本；
原始 .wsp 保持不動，供 compare_wsp.py 做一致性比對）。

修正紀錄：
- v2 改「依區段位置置換」（原 replace-by-content 會打到檔案中第一個
  相同內容的區段，例如 IndicatorHelper 的 Price input）
- v2 同名 input 出現於多個 SignalObject（多圖多訊號）時全部寫入
- v2 價格枚舉 input（.wsp 存 (h+l)*0.5 等運算式＝Median）不寫入，
  與 compare_wsp 的 SPECIAL 認定一致

用法：
    python scripts\\param_set.py --main <彙整 xlsx> --wsp-dir <E:\\Jobe\\群益交易程式>
                                 [--sids 0 2 5]   # 只處理指定策略編號（預設全部）
                                 [--out-suffix _BEST]
"""
from __future__ import annotations

import argparse
import io
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from compare_wsp import parse_report_params, DERIVED_COLS, PRICE_EXPRS  # noqa: E402

RE_INPUT_SECTION = re.compile(
    r"^\[Wsp\\(Window_\d+)\\ChartManager\\Strategy\\SignalObject_\d+"
    r"\\SignalHelper\\InputHelper\\(Input_[^\]]+)\]\s*\n(.*?)(?=\n\[|\Z)",
    re.DOTALL | re.MULTILINE,
)
RE_LINE = re.compile(r"(?m)^(\s*)(Value|LocalCurrentValue2)(\s*=\s*')([^']*)('.*)$")
RE_SCALAR = re.compile(r"(?m)^\s*Value\s*=\s*'([^']*)'\s*$")


def parse_wsp_inputs(path):
    """回傳 (blocks, text)。blocks: [{win, iname, block, s, e}]，s/e 為區段內容位置。"""
    text = open(path, "rb").read().decode("big5", "replace")
    blocks = []
    for m in RE_INPUT_SECTION.finditer(text):
        blocks.append({
            "win": m.group(1),
            "iname": m.group(2)[len("Input_"):],
            "block": m.group(3),
            "s": m.start(3),
            "e": m.end(3),
        })
    return blocks, text


def current_value(block):
    m = RE_SCALAR.search(block)
    return m.group(1) if m else None


def build_best_copy(wsp_path, best_params, out_path):
    """把 best_params {name: raw} 依位置寫入 .wsp 副本（Value + LocalCurrentValue2）。

    作法：先對每個區段算出 newblock（不改原 text），收集成 (s, e, newblock)，
    最後依位置「由後往前」套用（避免前面的 splice 造成後面偏移失效）。

    回傳 (changed, skipped, report_only, price_skipped)
    """
    blocks, text = parse_wsp_inputs(wsp_path)
    changed, skipped, report_only, price_skipped = [], [], [], []

    by_name = {}
    for b in blocks:
        by_name.setdefault(b["iname"].lower(), []).append(b)

    edits = []  # (s, e, newblock)
    for name, rv in sorted(best_params.items()):
        if name in DERIVED_COLS:
            continue
        targets = by_name.get(name.lower())
        if not targets:
            report_only.append(name)
            continue
        for b in targets:
            cur = current_value(b["block"])
            # 價格枚舉特例：.wsp 存運算式（如 (h+l)*0.5）＝報告枚舉值（0）同義，不動
            if cur is not None:
                wsp_key = re.sub(r"\s+", "", cur.lower())
                if wsp_key in PRICE_EXPRS:
                    price_skipped.append((b["win"], b["iname"], cur))
                    continue
            write_val = rv
            if cur is not None and cur.strip().lower() in ("true", "false"):
                low = rv.strip().lower()
                if low == "1":
                    write_val = "true"
                elif low == "0":
                    write_val = "false"
            newblock, n = RE_LINE.subn(
                lambda m: "%s%s%s%s%s" % (m.group(1), m.group(2), m.group(3),
                                          write_val, m.group(5)),
                b["block"],
            )
            if n == 0:
                changed.append((b["win"], name, cur, write_val, "no match"))
                continue
            edits.append((b["s"], b["e"], newblock))
            changed.append((b["win"], name, cur, write_val, n))

    # 由後往前套用
    for s, e, nb in sorted(edits, key=lambda x: -x[0]):
        text = text[:s] + nb + text[e:]

    for name in sorted(set(by_name) - set(k.lower() for k in best_params)):
        for b in by_name[name]:
            skipped.append((b["win"], b["iname"], current_value(b["block"])))

    with open(out_path, "wb") as f:
        f.write(text.encode("big5", "replace"))
    return changed, skipped, report_only, price_skipped


def find_wsp(sid, name, wsp_dir):
    cand = os.path.join(wsp_dir, "%s-%s.wsp" % (sid, re.sub(r"^\d+-", "", name)))
    if os.path.isfile(cand):
        return cand
    for f in os.listdir(wsp_dir):
        if f.lower().endswith(".wsp") and f.startswith("%s-" % sid):
            return os.path.join(wsp_dir, f)
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--main",
                    default=r"H:\我的雲端硬碟\最佳化參數彙整\策略最佳化參數彙整_20260928.xlsx")
    ap.add_argument("--wsp-dir", default=r"E:\Jobe\群益交易程式")
    ap.add_argument("--sids", nargs="*", default=None)
    ap.add_argument("--out-suffix", default="_BEST")
    ap.add_argument("--out-dir", default=None)
    args = ap.parse_args()

    import openpyxl
    wb = openpyxl.load_workbook(args.main)
    ws = wb["彙整主表"]
    out = io.open(r"C:\Users\Administrator\AppData\Local\Temp\opencode\param_set_report.txt",
                  "w", encoding="utf-8")
    made = 0
    for r in range(5, ws.max_row + 1):
        sid = str(ws.cell(row=r, column=1).value).strip()
        name = str(ws.cell(row=r, column=2).value).strip()
        p14 = ws.cell(row=r, column=14).value
        if args.sids and sid not in args.sids:
            continue
        if not p14:
            out.write("SID=%s %s: 主表無 col14，跳過\n" % (sid, name))
            continue
        wsp_path = find_wsp(sid, name, args.wsp_dir)
        if not wsp_path:
            out.write("SID=%s %s: 找不到 .wsp\n" % (sid, name))
            continue
        best = parse_report_params(p14)
        outdir = args.out_dir or os.path.dirname(wsp_path)
        out_path = os.path.join(
            outdir, os.path.basename(wsp_path).replace(".wsp", "%s.wsp" % args.out_suffix))
        changed, skipped, report_only, price_skipped = build_best_copy(wsp_path, best, out_path)
        made += 1
        out.write("SID=%s %s\n  wsp=%s\n" % (sid, name, os.path.basename(wsp_path)))
        grp = {}
        for win, nm, cur, newv, n in changed:
            grp.setdefault(nm, []).append((win, cur, newv, n))
        if grp:
            out.write("  寫入 %d 組參數：\n" % len(grp))
            for nm, items in sorted(grp.items()):
                for win, cur, newv, n in items:
                    out.write("    [%s] %-20s %s -> %s (改 %d 行)\n" % (win, nm, cur, newv, n))
        if price_skipped:
            out.write("  價格枚舉特例（未寫入，同義）：%s\n" %
                      ", ".join("%s.%s=%s" % (w, n, v) for w, n, v in price_skipped))
        if skipped:
            out.write("  wsp 有而最佳參數未包含（不變更）：%s\n" %
                      ", ".join("%s.%s=%s" % s for s in skipped))
        if report_only:
            out.write("  最佳參數有而 wsp 無對應 input：%s\n" % ", ".join(report_only))
        out.write("  out=%s\n" % out_path)
    out.close()
    print("made %d copies" % made)


if __name__ == "__main__":
    main()